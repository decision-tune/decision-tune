"""Recipes: the one unit every DecisionTune surface shares (Python, CLI, HTTP server, MCP, app).

A recipe says what to read from each item, which questions to ask, and when to flag a row for a person.

    from decision_tune import Recipe
    cols, rows = read_rows("tickets.csv")
    out = Recipe.load("support-triage").run(rows)       # one dict per row: its columns + the answers + needs_review
"""
import csv
import io
import json
import os
import re
import tempfile
from pathlib import Path

NAME_RE = re.compile(r"^[a-z0-9][a-z0-9_-]{0,63}$")
TYPES = ("choose", "yes_no", "scale")
BUILTIN = Path(__file__).parent / "recipes"
_model = None


def recipes_dir():
    home = os.environ.get("DECISION_TUNE_HOME")  # set but empty still counts as set
    return os.path.join(os.path.expanduser(home if home is not None else "~/.decision-tune"), "recipes")


def _name(v, what):
    if not isinstance(v, str) or not NAME_RE.fullmatch(v):
        raise ValueError(f"{what} name {v!r} must match {NAME_RE.pattern}")
    return v


def _choices(q, key):
    v = q.get(key)
    if isinstance(v, dict):
        ok = all(isinstance(k, str) and k and isinstance(d, str) for k, d in v.items())
    else:
        ok = isinstance(v, list) and all(isinstance(o, str) and o for o in v) and len(set(v)) == len(v)
    if not ok or not 2 <= len(v) <= 32:
        raise ValueError(f"question {q['name']!r}: {key} must be 2 to 32 entries, a {{key: description}} dict or a list of strings")
    return v


class Recipe:
    def __init__(self, name, read, questions, review_below):
        self.name, self.read, self.questions, self.review_below = name, read, questions, review_below

    @classmethod
    def from_dict(cls, d):
        if not isinstance(d, dict):
            raise ValueError("a recipe must be a JSON object")
        name = _name(d.get("name"), "recipe")
        read = d.get("read", [])
        if not isinstance(read, list) or not all(isinstance(c, str) and c for c in read):
            raise ValueError("read must be a list of column names")
        qs = d.get("questions")
        if not isinstance(qs, list) or not 1 <= len(qs) <= 50:
            raise ValueError("questions must be a list of 1 to 50 questions")
        out = []
        for q in qs:
            if not isinstance(q, dict):
                raise ValueError("each question must be an object")
            _name(q.get("name"), "question")
            if q.get("type") not in TYPES:
                raise ValueError(f"question {q['name']!r}: type must be one of {', '.join(TYPES)}")
            if not isinstance(q.get("question"), str) or not q["question"].strip():
                raise ValueError(f"question {q['name']!r}: question text is required")
            item = {"name": q["name"], "type": q["type"], "question": q["question"]}
            if q["type"] != "yes_no":
                key = "options" if q["type"] == "choose" else "levels"
                item[key] = _choices(q, key)
            out.append(item)
        if len({q["name"] for q in out}) < len(out):
            raise ValueError("question names must be unique")
        owner = {"needs_review": None, "error": None}  # output field -> question that owns it (None = reserved)
        for q in out:
            for f in (q["name"], f"{q['name']}_confidence", f"{q['name']}_p_yes"):  # p_yes reserved for every type
                if f in owner:
                    who = "a reserved field" if owner[f] is None else f"question {owner[f]!r}"
                    raise ValueError(f"question {q['name']!r}: output field {f!r} clashes with {who}")
                owner[f] = q["name"]
        rb = d.get("review_below", 0.6)
        if isinstance(rb, bool) or not isinstance(rb, (int, float)) or not 0 <= rb <= 1:
            raise ValueError("review_below must be a number from 0 to 1")
        return cls(name, list(read), out, float(rb))

    @classmethod
    def load(cls, name_or_path):
        p = os.path.expanduser(str(name_or_path))
        if not (p.endswith(".json") and os.path.isfile(p)):
            p = next((str(f) for f in (Path(recipes_dir(), f"{p}.json"), BUILTIN / f"{p}.json") if NAME_RE.fullmatch(p) and f.is_file()), None)
        if p is None:
            have = ", ".join(r["name"] for r in list_recipes()) or "none"
            raise ValueError(f"no recipe {str(name_or_path)!r}; available: {have}")
        try:
            with open(p, encoding="utf-8") as f:
                return cls.from_dict(json.load(f))
        except json.JSONDecodeError as e:
            raise ValueError(f"{p} is not valid JSON: {e}") from None

    def to_dict(self):
        return {"name": self.name, "read": list(self.read), "questions": json.loads(json.dumps(self.questions)),
                "review_below": self.review_below}

    def save(self, directory=None):
        _name(self.name, "recipe")  # the attribute is mutable and the constructor unchecked: recheck before touching disk
        d = directory or recipes_dir()
        os.makedirs(d, exist_ok=True)
        path = os.path.join(d, f"{self.name}.json")
        fd, tmp = tempfile.mkstemp(dir=d, prefix=f".{self.name}.", suffix=".tmp")  # unique, O_EXCL, never a followed symlink
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as f:
                json.dump(self.to_dict(), f, indent=2, ensure_ascii=False)
                f.write("\n")
            os.replace(tmp, path)  # ponytail: a user copy shadows a built-in of the same name; delete the file to unshadow
        except BaseException:
            if os.path.lexists(tmp):
                os.unlink(tmp)
            raise
        return path

    def state(self, row):
        cols = self.read or list(row)
        cell = lambda c: "" if row.get(c) is None else row[c]
        return str(cell(cols[0])) if len(cols) == 1 else {c: cell(c) for c in cols}  # one column = that cell's text

    def decide_row(self, model, row):
        """Answers for one row (dict of column -> value). A model error flags the row instead of raising."""
        out, conf = {}, []
        try:
            state = self.state(row)
            for q in self.questions:
                n = q["name"]
                if q["type"] == "yes_no":
                    p = model.yes_no(state, q["question"])
                    c = max(p, 1 - p)
                    out.update({n: "yes" if p >= 0.5 else "no", f"{n}_p_yes": round(p, 4), f"{n}_confidence": round(c, 4)})
                else:
                    r = model.choose(state, q["question"], q["options"] if q["type"] == "choose" else q["levels"])
                    c = r["confidence"]
                    out.update({n: r["choice"], f"{n}_confidence": round(c, 4)})
                conf.append(c)
        except Exception as e:  # e.g. a sequence over the context limit: keep the run going
            msg = str(e).splitlines()[0] if str(e) else ""
            return {"error": f"{type(e).__name__}: {msg}", "needs_review": True}
        out["needs_review"] = any(c < self.review_below for c in conf)
        return out

    def run(self, rows, model=None, progress=None):
        global _model
        if model is None:
            if _model is None:
                from .hub import DecisionModel

                _model = DecisionModel.from_pretrained()
            model = _model
        results = []
        for i, row in enumerate(rows, 1):
            results.append({**row, **self.decide_row(model, row)})
            if progress:
                progress(i, len(rows))
        return results

    def output_columns(self, input_columns):
        cols = list(input_columns)
        for q in self.questions:
            n = q["name"]
            cols += [n, f"{n}_p_yes", f"{n}_confidence"] if q["type"] == "yes_no" else [n, f"{n}_confidence"]
        return list(dict.fromkeys(cols + ["needs_review", "error"]))


def list_recipes():
    """User recipes (recipes_dir()) and built-in ones; a user recipe shadows a built-in of the same name."""
    found = {}
    for source, d in (("built-in", BUILTIN), ("user", Path(recipes_dir()))):
        for f in sorted(d.glob("*.json")) if d.is_dir() else []:
            try:
                r = Recipe.load(f) if source == "user" else Recipe.from_dict(json.loads(f.read_text(encoding="utf-8")))
            except ValueError:
                continue  # a broken file must not hide the others
            if r.name == f.stem:
                found[r.name] = {"name": r.name, "source": source, "questions": [q["name"] for q in r.questions], "path": str(f)}
    return sorted(found.values(), key=lambda x: x["name"])


def _rows(columns, records):
    return columns, [{c: ("" if r.get(c) is None else r[c]) for c in columns} for r in records]


def rows_from_csv_text(text):
    rd = csv.DictReader(io.StringIO(text.lstrip("﻿")))
    return _rows(list(rd.fieldnames or []), rd)


def rows_from_lines(text):
    return ["text"], [{"text": s} for s in (ln.strip() for ln in text.splitlines()) if s]


def read_rows(path):
    """(columns, rows) from a .csv, an .xlsx (first sheet, first row = header) or a folder of .txt/.md files."""
    p = Path(os.path.expanduser(str(path)))
    ext = p.suffix.lower()
    if p.is_dir():
        files = sorted(f for f in p.iterdir() if f.is_file() and f.suffix.lower() in (".txt", ".md"))
        return ["file", "text"], [{"file": f.name, "text": f.read_text(encoding="utf-8", errors="replace")} for f in files]
    if ext == ".csv":
        return rows_from_csv_text(p.read_text(encoding="utf-8-sig"))
    if ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            it = wb.worksheets[0].iter_rows(values_only=True)
            cols = [str(c) for c in next(it, ())]
            return _rows(cols, (dict(zip(cols, r)) for r in it if any(v is not None for v in r)))
        finally:
            wb.close()
    raise ValueError(f"cannot read {str(path)!r}: use a .csv, an .xlsx or a folder of .txt/.md files")


def safe_cell(v):
    """Spreadsheet formula-injection guard: text starting with = + - @ TAB or CR gets a leading apostrophe."""
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


def write_csv(path_or_file, columns, rows):
    def put(f):
        w = csv.writer(f)
        w.writerow([safe_cell(c) for c in columns])
        w.writerows([safe_cell(r.get(c, "")) for c in columns] for r in rows)

    if hasattr(path_or_file, "write"):
        return put(path_or_file)
    with open(path_or_file, "w", newline="", encoding="utf-8-sig") as f:  # BOM so Excel reads accents right
        put(f)


def write_xlsx(path, columns, rows):
    import openpyxl
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def cell(v):  # strip first: removing a control char can expose a leading = + - @
        return safe_cell(ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([cell(c) for c in columns])
    for r in rows:
        ws.append([cell(r.get(c, "")) for c in columns])
    wb.save(path)
