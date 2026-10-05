import csv
import io
import json
import os
import subprocess
import sys

import openpyxl
import pytest

from decision_tune import __version__
from decision_tune import mcp
from decision_tune.mcp import run_stdio

SRC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src")

DATA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "tickets.csv")
OPTS = {"billing": "Billing: charges", "shipping": "Shipping: delivery", "tech": "Tech support: bugs"}


class Fake:
    """Scripted answers: choose picks 'shipping', yes_no says 0.9."""

    def __init__(self):
        self.calls = []

    def choose(self, state, question, options):
        self.calls.append(("choose", state, question, options))
        keys = list(options)
        probs = {k: 0.1 for k in keys}
        probs["shipping" if "shipping" in keys else keys[0]] = 0.8
        best = max(probs, key=probs.get)
        return {"choice": best, "probabilities": probs, "confidence": probs[best]}

    def yes_no(self, state, question):
        self.calls.append(("yes_no", state, question))
        return 0.9


def strict(line):
    def no_const(c):
        raise ValueError(c)

    return json.loads(line, parse_constant=no_const)


class Scripted:
    """Different scripted result per call, so ignored or stale results show up."""

    def __init__(self, choices, yes):
        self.choices, self.yes, self.calls = list(choices), list(yes), []

    def choose(self, state, question, options):
        self.calls.append(("choose", state, question, options))
        c = self.choices.pop(0)
        return {"choice": c, "probabilities": {k: (0.7 if k == c else 0.1) for k in options}, "confidence": 0.7}

    def yes_no(self, state, question):
        self.calls.append(("yes_no", state, question))
        return self.yes.pop(0)


def rpc(*msgs, factory=None):
    """Send each message (dict or raw str) as one line; return (parsed response lines, raw stdout, factory call count)."""
    fake, n = Fake(), []

    def make():
        n.append(1)
        return (factory or (lambda: fake))()

    lines = "".join((m if isinstance(m, str) else json.dumps(m)) + "\n" for m in msgs)
    out = io.StringIO()
    run_stdio(make, io.StringIO(lines), out)
    return [strict(x) for x in out.getvalue().splitlines()], out.getvalue(), len(n)


def req(id_, method, **params):
    return {"jsonrpc": "2.0", "id": id_, "method": method, "params": params}


def call(id_, name, **args):
    return req(id_, "tools/call", name=name, arguments=args)


def one(msg):
    (r,), _, _ = rpc(msg)
    return r


def test_initialize_negotiates_version():
    r = one(req(1, "initialize", protocolVersion="2025-03-26", capabilities={}, clientInfo={"name": "t", "version": "1"}))
    res = r["result"]
    assert r["id"] == 1 and r["jsonrpc"] == "2.0"
    assert res["protocolVersion"] == "2025-03-26"
    assert res["capabilities"] == {"tools": {}}
    assert res["serverInfo"] == {"name": "decisiontune", "version": __version__}
    assert isinstance(res["instructions"], str) and res["instructions"]
    assert one(req(2, "initialize", protocolVersion="2099-01-01"))["result"]["protocolVersion"] == "2025-06-18"
    assert one(req(3, "initialize"))["result"]["protocolVersion"] == "2025-06-18"


def test_notification_gets_no_output():
    out, raw, _ = rpc({"jsonrpc": "2.0", "method": "notifications/initialized"})
    assert out == [] and raw == ""


def test_ping_and_blank_lines():
    out, _, _ = rpc("", req(1, "ping"), "   ")
    assert out == [{"jsonrpc": "2.0", "id": 1, "result": {}}]


def test_tools_list():
    tools = one(req(1, "tools/list"))["result"]["tools"]
    assert [t["name"] for t in tools] == ["decide", "run_recipe", "list_recipes"]
    for t in tools:
        assert t["description"] and t["inputSchema"]["type"] == "object"
    assert tools[0]["inputSchema"]["required"] == ["state", "question"]
    assert tools[1]["inputSchema"]["required"] == ["recipe"]


def test_decide_choose():
    res = one(call(1, "decide", state="The order arrived broken.", question="Which team?", options=OPTS))["result"]
    assert res["isError"] is False
    s = res["structuredContent"]
    assert s["choice"] == "shipping" and s["confidence"] == 0.8 and set(s["probabilities"]) == set(OPTS) and s["ms"] >= 0
    text = res["content"][0]
    assert text["type"] == "text" and "shipping" in text["text"] and json.loads(text["text"].split("\n", 1)[1]) == s


def test_decide_choose_list_options():
    s = one(call(1, "decide", state={"a": 1}, question="Which?", options=["x", "y"]))["result"]["structuredContent"]
    assert s["choice"] == "x"


def test_decide_yes_no():
    res = one(call(1, "decide", state="Please refund me.", question="Is it a refund?"))["result"]
    assert res["isError"] is False
    assert res["structuredContent"]["answer"] == "yes" and res["structuredContent"]["p_yes"] == 0.9 and "ms" in res["structuredContent"]


@pytest.mark.parametrize("args", [
    {"question": "q"},
    {"state": "s"},
    {"state": 5, "question": "q"},
    {"state": "s", "question": "q", "options": "billing"},
    {"state": "s", "question": "q", "options": [1, 2]},
])
def test_decide_bad_input_is_a_tool_error(args):
    res = one(call(1, "decide", **args))["result"]
    assert res["isError"] is True and "Traceback" not in res["content"][0]["text"] and "\n" not in res["content"][0]["text"]


def test_run_recipe_on_csv():
    res = one(call(1, "run_recipe", recipe="support-triage", input_path=DATA))["result"]
    assert res["isError"] is False
    s = res["structuredContent"]
    assert s["count"] == 8 and len(s["rows"]) == 8 and s["truncated"] is False and "output_path" not in s
    assert s["rows"][0]["team"] == "shipping" and s["rows"][0]["refund"] == "yes" and s["rows"][0]["tone"] in ("neg", "neu", "pos")
    assert isinstance(s["needs_review"], int) and s["ms"] >= 0


def test_run_recipe_with_rows_and_recipe_object():
    recipe = {"name": "t", "read": ["a"], "questions": [{"name": "q", "type": "yes_no", "question": "Is it?"}]}
    s = one(call(1, "run_recipe", recipe=recipe, rows=[{"a": "x"}, {"a": "y"}]))["result"]["structuredContent"]
    assert s["count"] == 2 and s["rows"][1] == {"a": "y", "q": "yes", "q_p_yes": 0.9, "q_confidence": 0.9, "needs_review": False}


def test_run_recipe_truncates_to_200_rows():
    recipe = {"name": "t", "read": ["a"], "questions": [{"name": "q", "type": "yes_no", "question": "Is it?"}]}
    s = one(call(1, "run_recipe", recipe=recipe, rows=[{"a": str(i)} for i in range(205)]))["result"]["structuredContent"]
    assert s["count"] == 205 and len(s["rows"]) == 200 and s["truncated"] is True


def test_run_recipe_writes_csv(tmp_path):
    out = tmp_path / "tickets-decided.csv"
    s = one(call(1, "run_recipe", recipe="support-triage", input_path=DATA, output_path=str(out)))["result"]["structuredContent"]
    assert s["output_path"] == str(out)
    got = list(csv.DictReader(open(out, encoding="utf-8-sig")))
    assert len(got) == 8 and got[0]["team"] == "shipping" and "needs_review" in got[0] and "team_confidence" in got[0]
    # a -decided file may be replaced, with the new content
    small = tmp_path / "small.csv"
    small.write_text("subject,message\nhi,Where is my order?\nyo,I was charged twice\nhey,The app crashes\n")
    res = one(call(1, "run_recipe", recipe="support-triage", input_path=str(small), output_path=str(out)))["result"]
    assert res["isError"] is False and res["structuredContent"]["output_path"] == str(out)
    got = list(csv.DictReader(open(out, encoding="utf-8-sig")))
    assert [r["subject"] for r in got] == ["hi", "yo", "hey"]
    assert [p.name for p in tmp_path.iterdir() if p.name.startswith(".")] == []  # no temp file left behind


def test_run_recipe_writes_xlsx_with_formula_guard(tmp_path):
    out = tmp_path / "x.xlsx"
    recipe = {"name": "t", "read": ["a"], "questions": [{"name": "q", "type": "yes_no", "question": "Is it?"}]}
    res = one(call(1, "run_recipe", recipe=recipe, rows=[{"a": "=1+1"}], output_path=str(out)))["result"]
    assert res["isError"] is False
    ws = openpyxl.load_workbook(out).active
    assert [c.value for c in ws[1]][:2] == ["a", "q"] and ws["A2"].value == "'=1+1"


def test_run_recipe_expands_tilde(tmp_path, monkeypatch):
    monkeypatch.setenv("HOME", str(tmp_path))
    (tmp_path / "t.csv").write_text("subject,message\nhi,Where is my order?\n")
    s = one(call(1, "run_recipe", recipe="support-triage", input_path="~/t.csv"))["result"]["structuredContent"]
    assert s["count"] == 1


def test_output_path_never_clobbers(tmp_path):
    keep = tmp_path / "mine.csv"
    keep.write_text("precious")
    link = tmp_path / "evil-decided.csv"
    link.symlink_to(keep)
    for bad in (keep, link, tmp_path / "no" / "such" / "x.csv", tmp_path / "x.txt"):
        res = one(call(1, "run_recipe", recipe="support-triage", input_path=DATA, output_path=str(bad)))["result"]
        assert res["isError"] is True, bad
    assert keep.read_text() == "precious"
    assert not (tmp_path / "x.txt").exists()


def test_output_path_checked_before_the_model_runs(tmp_path):
    keep = tmp_path / "mine.csv"
    keep.write_text("precious")
    _, _, loads = rpc(call(1, "run_recipe", recipe="support-triage", input_path=DATA, output_path=str(keep)))
    assert loads == 0


@pytest.mark.parametrize("args", [
    {"recipe": "no-such-recipe", "rows": []},
    {"recipe": "support-triage"},
    {"recipe": "support-triage", "rows": [], "input_path": DATA},
    {"recipe": "support-triage", "input_path": "/no/such/file.csv"},
    {"recipe": "support-triage", "input_path": DATA + ".txt"},
    {"recipe": "support-triage", "rows": ["not an object"]},
    {"recipe": {"name": "bad"}, "rows": []},
    {"recipe": 5, "rows": []},
])
def test_run_recipe_bad_input_is_a_tool_error(args):
    res = one(call(1, "run_recipe", **args))["result"]
    assert res["isError"] is True and "\n" not in res["content"][0]["text"] and "Traceback" not in res["content"][0]["text"]


def test_list_recipes(tmp_path, monkeypatch):
    monkeypatch.setenv("DECISION_TUNE_HOME", str(tmp_path))
    s = one(call(1, "list_recipes"))["result"]["structuredContent"]
    assert "support-triage" in [r["name"] for r in s["recipes"]]


def test_unknown_tool_and_method_and_bad_json():
    res = one(call(1, "nope"))["result"]
    assert res["isError"] is True and "nope" in res["content"][0]["text"]
    r = one(req(2, "resources/list"))
    assert r["id"] == 2 and r["error"]["code"] == -32601 and "result" not in r
    r = one("{not json")
    assert r["id"] is None and r["error"]["code"] == -32700
    assert one("[1, 2]")["error"]["code"] == -32600
    assert one({"jsonrpc": "2.0", "id": 4, "method": "tools/list", "params": "x"})["error"]["code"] == -32602


def test_every_request_gets_exactly_one_response_in_order():
    out, _, _ = rpc(req(1, "ping"), {"jsonrpc": "2.0", "method": "notifications/initialized"}, "{bad", req("b", "ping"),
                    call(3, "nope"))
    assert [r["id"] for r in out] == [1, None, "b", 3]


def test_model_loaded_lazily_and_used():
    model = Scripted(["billing", "tech"], [0.9, 0.2])
    n = []

    def make():
        n.append(1)
        return model

    head = [req(1, "initialize", protocolVersion="2025-06-18"), {"jsonrpc": "2.0", "method": "notifications/initialized"},
            req(2, "tools/list"), req(3, "ping"), call(4, "list_recipes")]
    _, _, loads = rpc(*head, factory=make)
    assert loads == 0 and model.calls == []
    o1, o2 = {"billing": "Billing: charges", "tech": "Tech: bugs"}, ["x", "y"]
    out, _, loads = rpc(*head[:3], call(4, "decide", state="S1", question="Q1", options=o1),
                        call(5, "decide", state={"k": "S2"}, question="Q2", options=o2),
                        call(6, "decide", state="S3", question="Q3"), call(7, "decide", state="S4", question="Q4"), factory=make)
    assert loads == 1
    assert model.calls == [("choose", "S1", "Q1", o1), ("choose", {"k": "S2"}, "Q2", o2), ("yes_no", "S3", "Q3"), ("yes_no", "S4", "Q4")]
    s = [r["result"]["structuredContent"] for r in out[-4:]]
    assert s[0]["choice"] == "billing" and s[1]["choice"] == "tech" and set(s[0]["probabilities"]) == set(o1)
    assert (s[2]["answer"], s[2]["p_yes"]) == ("yes", 0.9) and (s[3]["answer"], s[3]["p_yes"]) == ("no", 0.2)


def test_model_load_failure_is_a_tool_error_and_retried():
    n = []

    def make():
        n.append(1)
        if len(n) == 1:
            raise RuntimeError("DecisionTune 1.0 is not downloaded yet.\nRun the download command.")
        return Fake()

    out, _, loads = rpc(call(1, "decide", state="s", question="q"), call(2, "decide", state="s", question="q"), req(3, "ping"), factory=make)
    assert [r["id"] for r in out] == [1, 2, 3]
    assert out[0]["result"]["isError"] is True and out[0]["result"]["content"][0]["text"] == "DecisionTune 1.0 is not downloaded yet."
    assert loads == 2 and len(n) == 2
    assert out[1]["result"]["isError"] is False and out[1]["result"]["structuredContent"]["answer"] == "yes"
    assert out[2]["result"] == {}


def test_stdout_is_protocol_only(capsys):
    class Noisy(Fake):
        pass

    def make():
        print("downloading the model...")  # a library printing to stdout must not reach the protocol stream
        return Noisy()

    msgs = [req(1, "initialize", protocolVersion="2025-06-18"), {"jsonrpc": "2.0", "method": "notifications/initialized"},
            req(2, "tools/list"), call(3, "decide", state="s", question="q", options=OPTS), call(4, "run_recipe", recipe="support-triage", input_path=DATA),
            call(5, "list_recipes"), call(6, "nope"), req(7, "bogus"), "{bad json", req(8, "ping")]
    _, raw, _ = rpc(*msgs, factory=make)
    lines = raw.splitlines()
    assert len(lines) == 9 and raw.endswith("\n")
    for line in lines:
        m = json.loads(line)
        assert m["jsonrpc"] == "2.0" and "id" in m and ("result" in m) != ("error" in m)
    cap = capsys.readouterr()
    assert cap.out == "" and "downloading the model..." in cap.err
    assert sys.stdout is not None and not sys.stdout.closed


def test_utf8_roundtrip():
    s = one(call(1, "decide", state="Où est mon colis ? 📦", question="Is it about delivery?"))["result"]["structuredContent"]
    assert s["answer"] == "yes"


# ---- review fixes (T3-fix1) ----

@pytest.mark.parametrize("params", [
    {"name": [], "arguments": {}},
    {"name": {}, "arguments": {}},
    {"name": 5, "arguments": {}},
    {"arguments": {}},
    {"name": "list_recipes", "arguments": []},
    {"name": "list_recipes", "arguments": False},
    {"name": "list_recipes", "arguments": ""},
    {"name": "list_recipes", "arguments": 0},
    {"name": "list_recipes", "arguments": "x"},
])
def test_tools_call_bad_shapes_are_tool_errors_not_internal_errors(params):
    r = one(req(1, "tools/call", **params))
    assert "error" not in r and r["result"]["isError"] is True and "\n" not in r["result"]["content"][0]["text"]


def test_tools_call_arguments_may_be_absent_or_null():
    assert one(req(1, "tools/call", name="list_recipes"))["result"]["isError"] is False
    assert one(req(2, "tools/call", name="list_recipes", arguments=None))["result"]["isError"] is False


@pytest.mark.parametrize("options", [{"x": 123}, {"x": None}, {"x": ["a"]}, {"x": {"y": "z"}}, {"a": "ok", "b": 1.5}])
def test_decide_option_descriptions_must_be_strings(options):
    out, _, loads = rpc(call(1, "decide", state="s", question="q", options=options))
    assert out[0]["result"]["isError"] is True and loads == 0


@pytest.mark.parametrize("bad", [
    {"jsonrpc": "1.0", "id": 1, "method": "ping"},
    {"id": 1, "method": "ping"},
    {"jsonrpc": 2.0, "id": 1, "method": "ping"},
    {"jsonrpc": "2.0", "id": {"a": 1}, "method": "ping"},
    {"jsonrpc": "2.0", "id": [1], "method": "ping"},
    {"jsonrpc": "2.0", "id": True, "method": "ping"},
    {"jsonrpc": "2.0", "id": 1.5, "method": "ping"},
    {"jsonrpc": "2.0", "id": None, "method": "ping"},
])
def test_bad_envelope_is_invalid_request(bad):
    r = one(bad)
    good_id = bad["id"] if isinstance(bad.get("id"), (str, int)) and not isinstance(bad["id"], bool) else None
    assert r["error"]["code"] == -32600 and "result" not in r and r["id"] == good_id


def test_good_ids_still_work():
    out, _, _ = rpc(req(0, "ping"), req("a", "ping"), req(-7, "ping"))
    assert [r["id"] for r in out] == [0, "a", -7]


@pytest.mark.parametrize("line", [
    '{"jsonrpc":"2.0","id":NaN,"method":"ping"}',
    '{"jsonrpc":"2.0","id":Infinity,"method":"ping"}',
    '{"jsonrpc":"2.0","id":1,"method":"tools/call","params":{"name":"decide","arguments":{"state":-Infinity,"question":"q"}}}',
])
def test_nonstandard_numbers_are_parse_errors(line):
    out, raw, _ = rpc(line, req(2, "ping"))
    assert out[0]["error"]["code"] == -32700 and out[0]["id"] is None and out[1]["result"] == {}
    assert "NaN" not in raw and "Infinity" not in raw


def test_nonfinite_values_never_reach_the_output():
    recipe = {"name": "t", "read": ["a"], "questions": [{"name": "q", "type": "yes_no", "question": "Is it?"}]}
    line = '{"jsonrpc":"2.0","id":1,"method":"tools/call","params":{"name":"run_recipe","arguments":{"recipe":%s,"rows":[{"a":1e999}]}}}' % json.dumps(recipe)
    out, raw, _ = rpc(line, req(2, "ping"))  # rpc parses strictly: NaN or Infinity in the output would raise
    assert [r["id"] for r in out] == [1, 2] and out[0]["result"]["isError"] is True


def test_deeply_nested_json_does_not_kill_the_server():
    out, _, _ = rpc("[" * 10000 + "]" * 10000, '{"a":' * 10000 + "1" + "}" * 10000, req(2, "ping"))
    assert [r.get("error", {}).get("code") for r in out[:2]] == [-32700, -32700] and out[2]["result"] == {}


def test_oversized_line_is_rejected_and_the_loop_goes_on(monkeypatch):
    monkeypatch.setattr(mcp, "MAX_LINE_BYTES", 200)
    big = json.dumps(req(1, "tools/call", name="decide", arguments={"state": "x" * 1000, "question": "q"}))
    out, _, loads = rpc(req(0, "ping"), big, req(2, "ping"))
    assert [r["id"] for r in out] == [0, None, 2] and out[1]["error"]["code"] == -32600 and out[2]["result"] == {} and loads == 0


def test_line_exactly_at_the_limit_is_accepted(monkeypatch):
    line = json.dumps(req(1, "ping"))
    monkeypatch.setattr(mcp, "MAX_LINE_BYTES", len(line.encode()))
    assert one(line)["result"] == {}


def test_input_row_and_workload_bounds(monkeypatch, tmp_path):
    recipe = {"name": "t", "read": ["a"], "questions": [{"name": "q", "type": "yes_no", "question": "Is it?"},
                                                         {"name": "r", "type": "yes_no", "question": "Is it?"}]}
    monkeypatch.setattr(mcp, "MAX_ROWS", 3)
    out, _, loads = rpc(call(1, "run_recipe", recipe=recipe, rows=[{"a": str(i)} for i in range(4)]),
                        call(2, "run_recipe", recipe=recipe, rows=[{"a": str(i)} for i in range(3)]))
    assert out[0]["result"]["isError"] is True and out[1]["result"]["isError"] is False and loads == 1
    f = tmp_path / "big.csv"
    f.write_text("a\n" + "\n".join("x" * 4 for _ in range(4)) + "\n")
    out, _, loads = rpc(call(1, "run_recipe", recipe=recipe, input_path=str(f)))
    assert out[0]["result"]["isError"] is True and loads == 0
    monkeypatch.setattr(mcp, "MAX_ROWS", 100)
    monkeypatch.setattr(mcp, "MAX_DECISIONS", 5)  # 3 rows x 2 questions = 6
    out, _, loads = rpc(call(1, "run_recipe", recipe=recipe, rows=[{"a": "1"}, {"a": "2"}, {"a": "3"}]),
                        call(2, "run_recipe", recipe=recipe, rows=[{"a": "1"}, {"a": "2"}]))
    assert out[0]["result"]["isError"] is True and out[1]["result"]["isError"] is False and loads == 1


def test_input_file_size_bound(monkeypatch, tmp_path):
    monkeypatch.setattr(mcp, "MAX_INPUT_BYTES", 10)
    out, _, loads = rpc(call(1, "run_recipe", recipe="support-triage", input_path=DATA))
    assert out[0]["result"]["isError"] is True and loads == 0


def test_special_files_are_rejected_not_read(tmp_path):
    fifo = tmp_path / "input.csv"
    os.mkfifo(fifo)  # a read would block forever
    for p in (fifo, "/dev/null"):
        res = one(call(1, "run_recipe", recipe="support-triage", input_path=str(p)))["result"]
        assert res["isError"] is True and "\n" not in res["content"][0]["text"]
    d = tmp_path / "folder"
    d.mkdir()
    (d / "a.txt").write_text("Where is my order?")
    assert one(call(2, "run_recipe", recipe="support-triage", input_path=str(d)))["result"]["isError"] is False


def test_destination_created_during_the_run_is_not_clobbered(tmp_path):
    dest = tmp_path / "export.csv"

    def make():  # runs after validation, before the write: the long race window
        dest.write_text("precious")
        return Fake()

    res = one_with(make, call(1, "run_recipe", recipe="support-triage", input_path=DATA, output_path=str(dest)))["result"]
    assert res["isError"] is True and dest.read_text() == "precious"
    assert [p.name for p in tmp_path.iterdir()] == ["export.csv"]  # temp file removed


def test_symlink_planted_during_the_run_is_not_written_through(tmp_path):
    keep = tmp_path / "mine.csv"
    keep.write_text("precious")
    for name in ("export.csv", "export-decided.csv"):
        dest = tmp_path / name

        def make():
            dest.symlink_to(keep)
            return Fake()

        res = one_with(make, call(1, "run_recipe", recipe="support-triage", input_path=DATA, output_path=str(dest)))["result"]
        assert res["isError"] is True and keep.read_text() == "precious"
        dest.unlink()
    assert sorted(p.name for p in tmp_path.iterdir()) == ["mine.csv"]


def test_hard_link_is_replaced_not_truncated(tmp_path):
    precious = tmp_path / "precious.csv"
    precious.write_text("precious")
    dest = tmp_path / "export-decided.csv"
    os.link(precious, dest)
    res = one(call(1, "run_recipe", recipe="support-triage", input_path=DATA, output_path=str(dest)))["result"]
    assert res["isError"] is False
    assert precious.read_text() == "precious"
    assert len(list(csv.DictReader(open(dest, encoding="utf-8-sig")))) == 8 and dest.stat().st_nlink == 1


def test_hard_link_with_a_disallowed_name_is_refused(tmp_path):
    precious = tmp_path / "precious.csv"
    precious.write_text("precious")
    dest = tmp_path / "export.csv"
    os.link(precious, dest)
    assert one(call(1, "run_recipe", recipe="support-triage", input_path=DATA, output_path=str(dest)))["result"]["isError"] is True
    assert precious.read_text() == dest.read_text() == "precious"


def test_new_output_file_has_normal_permissions(tmp_path):
    dest = tmp_path / "new.csv"
    one(call(1, "run_recipe", recipe="support-triage", input_path=DATA, output_path=str(dest)))
    umask = os.umask(0)
    os.umask(umask)
    assert dest.stat().st_mode & 0o777 == 0o666 & ~umask


def test_write_failure_leaves_no_temp_file(tmp_path, monkeypatch):
    def boom(*a):
        raise OSError("disk full")

    monkeypatch.setattr(mcp, "write_csv", boom)
    res = one(call(1, "run_recipe", recipe="support-triage", input_path=DATA, output_path=str(tmp_path / "a-decided.csv")))["result"]
    assert res["isError"] is True and list(tmp_path.iterdir()) == []


def one_with(factory, msg):
    (r,), _, _ = rpc(msg, factory=factory)
    return r


def test_native_fd1_writes_do_not_corrupt_the_protocol():
    code = (
        "import os\n"
        "from decision_tune.mcp import run_stdio\n"
        "class M:\n"
        "    def choose(self, s, q, o): return {'choice': list(o)[0], 'probabilities': {k: 0.5 for k in o}, 'confidence': 0.5}\n"
        "    def yes_no(self, s, q): return 0.9\n"
        "def make():\n"
        "    os.write(1, b'noise from a native library\\n')\n"
        "    return M()\n"
        "run_stdio(make)\n")
    msgs = [req(1, "initialize", protocolVersion="2025-06-18"), call(2, "decide", state="s", question="q"), req(3, "ping")]
    p = subprocess.run([sys.executable, "-c", code], input="".join(json.dumps(m) + "\n" for m in msgs), capture_output=True, text=True,
                       env={**os.environ, "PYTHONPATH": SRC}, timeout=120)
    lines = p.stdout.splitlines()
    assert len(lines) == 3, p.stdout + p.stderr
    assert [strict(x)["id"] for x in lines] == [1, 2, 3] and strict(lines[1])["result"]["structuredContent"]["answer"] == "yes"
    assert "noise from a native library" in p.stderr and "noise" not in p.stdout
