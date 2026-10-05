import csv
import io
import json
import os
import threading

import openpyxl
import pytest

from decision_tune import Recipe, list_recipes
from decision_tune.recipe import (read_rows, recipes_dir, rows_from_csv_text, rows_from_lines, safe_cell, write_csv,
                                  write_xlsx)

DATA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "tickets.csv")


def spec(**kw):
    d = {"name": "t", "read": ["a"], "questions": [{"name": "q", "type": "yes_no", "question": "Is it?"}]}
    d.update(kw)
    return d


@pytest.fixture
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("DECISION_TUNE_HOME", str(tmp_path))
    return tmp_path


class Fake:
    """Records every call; answers are scripted per question text."""

    def __init__(self, choose=None, yes=None):
        self.calls, self._c, self._y = [], choose or {}, yes or {}

    def choose(self, state, question, options):
        self.calls.append(("choose", state, question, options))
        return self._c(len(self.calls)) if callable(self._c) else self._c[question]

    def yes_no(self, state, question):
        self.calls.append(("yes_no", state, question))
        return self._y(len(self.calls)) if callable(self._y) else self._y[question]


def res(choice, conf):
    return {"choice": choice, "probabilities": {choice: conf}, "confidence": conf}


# 1. validation, load, save, list, readers, writers
@pytest.mark.parametrize("bad,msg", [
    (spec(name="Bad Name"), "name"),
    (spec(name="../x"), "name"),
    (spec(questions=[{"name": "q", "type": "rank", "question": "?"}]), "type"),
    (spec(questions=[{"name": "q", "type": "choose", "question": "?", "options": ["only"]}]), "options"),
    (spec(questions=[{"name": "q", "type": "scale", "question": "?", "levels": {}}]), "levels"),
    (spec(questions=[]), "questions"),
    (spec(questions=[{"name": "q", "type": "yes_no", "question": "?"}] * 2), "unique"),
    (spec(review_below=1.5), "review_below"),
    (spec(name="t\n"), "name"),  # 8: fullmatch, not match
    (spec(questions=[{"name": "q\n", "type": "yes_no", "question": "?"}]), "name"),
    (spec(read=False), "read"),  # 2: wrong types are not "all columns"
    (spec(read=0), "read"),
    (spec(read={}), "read"),
    (spec(read=None), "read"),
    (spec(read="a"), "read"),
    # 1 + accept 2: the whole output namespace, reserved fields included
    (spec(questions=[{"name": "q", "type": "yes_no", "question": "?"},
                     {"name": "q_confidence", "type": "yes_no", "question": "?"}]), "clashes"),
    (spec(questions=[{"name": "q", "type": "yes_no", "question": "?"},
                     {"name": "q_p_yes", "type": "yes_no", "question": "?"}]), "clashes"),
    (spec(questions=[{"name": "q", "type": "choose", "question": "?", "options": ["a", "b"]},
                     {"name": "q_confidence", "type": "yes_no", "question": "?"}]), "clashes"),
    (spec(questions=[{"name": "needs_review", "type": "yes_no", "question": "?"}]), "needs_review"),
    (spec(questions=[{"name": "error", "type": "yes_no", "question": "?"}]), "error"),
])
def test_validation_errors(bad, msg):
    with pytest.raises(ValueError, match=msg):
        Recipe.from_dict(bad)


def test_load_builtin_and_path(home, tmp_path):
    r = Recipe.load("support-triage")
    assert [q["name"] for q in r.questions] == ["team", "refund", "tone"] and r.review_below == 0.6
    p = tmp_path / "mine.json"
    p.write_text(json.dumps(spec(name="mine")))
    assert Recipe.load(str(p)).name == "mine"
    with pytest.raises(ValueError, match="support-triage"):
        Recipe.load("nope")
    with pytest.raises(ValueError):
        Recipe.load("../x")


def test_save_and_list_user_shadows_builtin(home):
    assert recipes_dir() == str(home / "recipes")
    assert {(x["name"], x["source"]) for x in list_recipes()} == {("support-triage", "built-in")}
    d = Recipe.from_dict(spec(name="support-triage", review_below=0.9)).save()
    assert d == str(home / "recipes" / "support-triage.json")
    Recipe.from_dict(spec(name="other")).save()
    got = {x["name"]: x for x in list_recipes()}
    assert got["support-triage"]["source"] == "user" and got["support-triage"]["questions"] == ["q"]
    assert got["other"]["source"] == "user" and len(got) == 2
    assert Recipe.load("support-triage").review_below == 0.9
    assert Recipe.from_dict(Recipe.load("other").to_dict()).to_dict() == Recipe.load("other").to_dict()


def test_read_rows_csv_xlsx_folder(tmp_path):
    cols, rows = read_rows(DATA)
    assert cols == ["subject", "message"] and len(rows) == 8 and rows[4]["message"].startswith("Thanks, the")
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    wb.active.append(["x", 1])
    wb.save(tmp_path / "t.xlsx")
    cols, rows = read_rows(str(tmp_path / "t.xlsx"))
    assert cols == ["a", "b"] and rows == [{"a": "x", "b": 1}]
    d = tmp_path / "notes"
    d.mkdir()
    (d / "b.md").write_text("two")
    (d / "a.txt").write_text("one")
    (d / "skip.pdf").write_text("no")
    assert read_rows(str(d)) == (["file", "text"], [{"file": "a.txt", "text": "one"}, {"file": "b.md", "text": "two"}])
    with pytest.raises(ValueError):
        read_rows(str(d / "skip.pdf"))
    assert rows_from_lines("a\n\n b \n") == (["text"], [{"text": "a"}, {"text": "b"}])
    assert rows_from_csv_text("x,y\n1,2\n") == (["x", "y"], [{"x": "1", "y": "2"}])


def test_formula_guard_csv_and_xlsx(tmp_path):
    assert [safe_cell(v) for v in ("=1+1", "+1", "-1", "@a", "\tx", "\rx", "ok", 5)] == ["'=1+1", "'+1", "'-1", "'@a", "'\tx", "'\rx", "ok", 5]
    rows = [{"a": "=HYPERLINK(1)", "b": "fine"}]
    buf = io.StringIO()
    write_csv(buf, ["a", "b"], rows)
    assert list(csv.reader(io.StringIO(buf.getvalue()))) == [["a", "b"], ["'=HYPERLINK(1)", "fine"]]
    write_csv(str(tmp_path / "o.csv"), ["a", "b"], rows)
    assert (tmp_path / "o.csv").read_text().splitlines()[1].startswith("'=")
    write_xlsx(str(tmp_path / "o.xlsx"), ["a", "b"], rows)
    ws = openpyxl.load_workbook(tmp_path / "o.xlsx").active
    assert [c.value for c in ws[2]] == ["'=HYPERLINK(1)", "fine"] and ws["A2"].data_type == "s"


# 2. PIN: the model is consulted for every question of every row
class PerRow(Fake):
    """Distinct scripted answers per row, keyed on the row's subject, so reuse of another row's answers shows up."""

    def __init__(self, subjects):
        super().__init__()
        self.idx = {s: i for i, s in enumerate(subjects)}

    def choose(self, state, question, options):
        self.calls.append(("choose", state, question, options))
        i, keys = self.idx[state["subject"]], list(options)
        return res(keys[i % len(keys)] if "team" in question else keys[-1 - i % len(keys)], 0.59 if i == 3 else 0.7 + i / 100)

    def yes_no(self, state, question):
        self.calls.append(("yes_no", state, question))
        return 0.62 + 0.04 * self.idx[state["subject"]]


def test_run_calls_model_per_question_per_row(home):
    cols, rows = read_rows(DATA)
    r = Recipe.load("support-triage")
    fake = PerRow([x["subject"] for x in rows])
    seen = []
    out = r.run(rows, model=fake, progress=lambda i, n: seen.append((i, n)))
    assert len(out) == 8 and len(fake.calls) == 24
    assert seen == [(i, 8) for i in range(1, 9)]
    team, refund, tone = r.questions
    for i, row in enumerate(rows):
        state = {"subject": row["subject"], "message": row["message"]}
        got = fake.calls[3 * i:3 * i + 3]
        assert [(c[0], c[1], c[2]) for c in got] == [("choose", state, team["question"]), ("yes_no", state, refund["question"]),
                                                      ("choose", state, tone["question"])]
        assert got[0][3] == team["options"] and got[2][3] == tone["levels"]
        t, k = list(team["options"])[i % 3], list(tone["levels"])[-1 - i % 3]
        c = 0.59 if i == 3 else round(0.7 + i / 100, 4)
        p = round(0.62 + 0.04 * i, 4)
        assert out[i] == {**row, "team": t, "team_confidence": c, "refund": "yes", "refund_p_yes": p, "refund_confidence": p,
                          "tone": k, "tone_confidence": c, "needs_review": i == 3}
    assert r.output_columns(cols) == ["subject", "message", "team", "team_confidence", "refund", "refund_p_yes",
                                      "refund_confidence", "tone", "tone_confidence", "needs_review", "error"]


def test_single_read_column_passes_text_and_empty_read_passes_all():
    one = Recipe.from_dict(spec(read=["a"]))
    f = Fake(yes={"Is it?": 0.9})
    one.decide_row(f, {"a": "hello", "b": "no"})
    assert f.calls[0][1] == "hello"
    allc = Recipe.from_dict(spec(read=[]))
    allc.decide_row(f, {"a": "1", "b": "2"})
    assert f.calls[1][1] == {"a": "1", "b": "2"}


# 3. PIN: review threshold is strict "below"
def test_needs_review_threshold():
    r = Recipe.from_dict({"name": "t", "read": ["a"], "review_below": 0.6, "questions": [
        {"name": "c", "type": "choose", "question": "C?", "options": ["x", "y"]},
        {"name": "y", "type": "yes_no", "question": "Y?"}]})
    def go(conf, p):
        return r.decide_row(Fake(choose={"C?": res("x", conf)}, yes={"Y?": p}), {"a": "t"})
    assert go(0.59, 0.9)["needs_review"] is True
    assert go(0.61, 0.9)["needs_review"] is False
    d = go(0.95, 0.45)
    assert d["needs_review"] is True and d["y"] == "no" and d["y_confidence"] == 0.55 and d["y_p_yes"] == 0.45
    # equality is NOT below; the comparison uses the raw value, not the 4-place rounded one
    assert go(0.6, 0.9)["needs_review"] is False
    assert go(0.59996, 0.9)["needs_review"] is True and go(0.59996, 0.9)["c_confidence"] == 0.6
    assert go(0.60004, 0.9)["needs_review"] is False
    assert go(0.95, 0.6)["needs_review"] is False and go(0.95, 0.4)["y_confidence"] == 0.6
    assert go(0.95, 0.59996)["needs_review"] is True and go(0.95, 0.40004)["needs_review"] is True
    assert go(0.95, 0.60004)["needs_review"] is False and go(0.95, 0.39996)["needs_review"] is False
    # the p_yes == 0.5 decision boundary: "yes", confidence 0.5
    d = go(0.95, 0.5)
    assert d["y"] == "yes" and d["y_confidence"] == 0.5 and d["needs_review"] is True
    r.review_below = 0.5
    assert go(0.95, 0.5)["needs_review"] is False and go(0.95, 0.4999)["y"] == "no"


def test_model_error_flags_row_instead_of_crashing():
    class Boom(Fake):
        def yes_no(self, state, question):
            raise ValueError("sequence of 9000 tokens exceeds the 8192-token context limit")
    out = Recipe.from_dict(spec()).run([{"a": "x"}, {"a": "y"}], model=Boom())
    assert len(out) == 2 and out[0]["needs_review"] is True and "9000 tokens" in out[0]["error"]


# fix1 regression tests
def test_output_namespace_clash_names_the_field():
    with pytest.raises(ValueError, match="'q_confidence'.*clashes"):
        Recipe.from_dict(spec(questions=[{"name": "q", "type": "yes_no", "question": "?"},
                                         {"name": "q_confidence", "type": "yes_no", "question": "?"}]))
    with pytest.raises(ValueError, match="'needs_review'.*reserved"):
        Recipe.from_dict(spec(questions=[{"name": "needs_review", "type": "yes_no", "question": "?"}]))
    Recipe.from_dict(spec(questions=[{"name": "q", "type": "yes_no", "question": "?"},
                                     {"name": "other", "type": "yes_no", "question": "?"}]))


def test_read_missing_means_all_columns():
    d = spec()
    del d["read"]
    assert Recipe.from_dict(d).read == []


def test_single_column_state_is_text():
    r, f = Recipe.from_dict(spec(read=["a"])), Fake(yes={"Is it?": 0.9})
    r.decide_row(f, {"a": 42})
    r.decide_row(f, {"a": 4.5})
    r.decide_row(f, {"a": None})
    assert [c[1] for c in f.calls] == ["42", "4.5", ""]
    r.decide_row(f, {"a": True})
    assert f.calls[-1][1] == "True"


def test_empty_home_is_set_not_unset(monkeypatch, tmp_path):
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.delenv("DECISION_TUNE_HOME", raising=False)
    assert recipes_dir() == str(tmp_path / ".decision-tune" / "recipes")
    monkeypatch.setenv("DECISION_TUNE_HOME", "")
    assert recipes_dir() == os.path.join("", "recipes") != str(tmp_path / ".decision-tune" / "recipes")


def test_xlsx_strips_illegal_chars_before_formula_guard(tmp_path):
    write_xlsx(str(tmp_path / "o.xlsx"), ["\x01=H", "b"], [{"\x01=H": "\x01=1+1", "b": "\x02\x00@x"}])
    ws = openpyxl.load_workbook(tmp_path / "o.xlsx").active
    assert [c.value for c in ws[1]] == ["'=H", "b"] and ws["A1"].data_type == "s"
    assert [c.value for c in ws[2]] == ["'=1+1", "'@x"]
    assert ws["A2"].data_type == "s" and ws["B2"].data_type == "s"


def test_save_revalidates_name_at_the_filesystem_boundary(tmp_path):
    out = tmp_path / "out"
    r = Recipe("../outside", ["a"], spec()["questions"], 0.6)
    with pytest.raises(ValueError, match="name"):
        r.save(str(out))
    ok = Recipe.from_dict(spec())
    ok.name = "../outside"
    with pytest.raises(ValueError, match="name"):
        ok.save(str(out))
    assert not (tmp_path / "outside.json").exists() and not out.exists()


def test_save_uses_unique_exclusive_temp_file(tmp_path):
    out, victim = tmp_path / "out", tmp_path / "victim.txt"
    out.mkdir()
    victim.write_text("keep")
    os.symlink(victim, out / "t.json.tmp")  # a planted predictable temp name must not redirect the write
    path = Recipe.from_dict(spec()).save(str(out))
    assert victim.read_text() == "keep" and json.load(open(path))["name"] == "t"
    assert sorted(os.listdir(out)) == ["t.json", "t.json.tmp"]  # only the planted link is left, no stray temp file


def test_concurrent_saves_do_not_collide(tmp_path):
    r, errs, gate = Recipe.from_dict(spec()), [], threading.Barrier(8)

    def go():
        gate.wait()
        try:
            for _ in range(30):
                r.save(str(tmp_path))
        except Exception as e:
            errs.append(e)

    ts = [threading.Thread(target=go) for _ in range(8)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert errs == [] and json.load(open(tmp_path / "t.json"))["name"] == "t" and os.listdir(tmp_path) == ["t.json"]


def test_load_rejects_trailing_newline_name(home):
    (home / "recipes").mkdir()
    (home / "recipes" / "t\n.json").write_text(json.dumps(spec()))
    with pytest.raises(ValueError):
        Recipe.load("t\n")
